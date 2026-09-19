from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path(r"D:\qimoruanj\DietApp\膳食营养分析系统汇报PPT.pptx")


COLORS = {
    "bg": RGBColor(247, 244, 238),
    "surface": RGBColor(255, 252, 248),
    "surface_alt": RGBColor(241, 247, 239),
    "primary": RGBColor(110, 167, 122),
    "primary_dark": RGBColor(54, 91, 64),
    "secondary": RGBColor(226, 143, 96),
    "teal": RGBColor(99, 170, 162),
    "text": RGBColor(36, 49, 38),
    "muted": RGBColor(103, 112, 102),
    "border": RGBColor(220, 212, 200),
    "white": RGBColor(255, 255, 255),
}


def slide_size(slide):
    prs = slide.part.package.presentation_part.presentation
    return prs.slide_width, prs.slide_height


def set_run_style(run, size, color, bold=False):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def set_text_style(paragraph, size, color, bold=False, level=0, align=PP_ALIGN.LEFT):
    paragraph.level = level
    paragraph.alignment = align
    if paragraph.runs:
        for run in paragraph.runs:
            set_run_style(run, size, color, bold)


def add_full_background(slide, color):
    width, height = slide_size(slide)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_canvas(slide):
    width, height = slide_size(slide)
    add_full_background(slide, COLORS["bg"])
    circle = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, width - Inches(2.2), -Inches(0.7), Inches(2.5), Inches(2.5)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = COLORS["surface_alt"]
    circle.line.fill.background()
    blob = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.2), height - Inches(1.0), Inches(2.0), Inches(0.7)
    )
    blob.fill.solid()
    blob.fill.fore_color.rgb = COLORS["surface_alt"]
    blob.line.fill.background()


def add_header(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(8.8), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    set_text_style(p, 26, COLORS["text"], bold=True)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.78), Inches(1.15), Inches(8.8), Inches(0.4))
        tf = sub_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = subtitle
        set_text_style(p, 11, COLORS["muted"])


def add_footer(slide, page_num):
    box = slide.shapes.add_textbox(Inches(11.95), Inches(7.0), Inches(0.6), Inches(0.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = str(page_num)
    set_text_style(p, 10, COLORS["muted"], align=PP_ALIGN.RIGHT)


def add_bullet_block(slide, left, top, width, height, title, bullets):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["surface"]
    card.line.color.rgb = COLORS["border"]

    title_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), Inches(0.4))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    set_text_style(p, 16, COLORS["primary_dark"], bold=True)

    text_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.6), width - Inches(0.45), height - Inches(0.8))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()

    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.bullet = True
        set_text_style(p, 13, COLORS["text"])


def add_tag(slide, text, left, top, fill):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, Inches(1.45), Inches(0.35)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    box = slide.shapes.add_textbox(left, top + Inches(0.02), Inches(1.45), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    set_text_style(p, 10, COLORS["primary_dark"], bold=True, align=PP_ALIGN.CENTER)


def add_card(slide, left, top, width, height, title, body, accent):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["surface"]
    shape.line.color.rgb = COLORS["border"]

    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + Inches(0.18), top + Inches(0.18), Inches(0.42), Inches(0.14)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.42), width - Inches(0.36), Inches(0.45))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    set_text_style(p, 16, COLORS["text"], bold=True)

    body_box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.9), width - Inches(0.36), height - Inches(1.05))
    tf = body_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    set_text_style(p, 12, COLORS["muted"])


def build_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(0.9), Inches(5.1), Inches(0.45)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = COLORS["surface_alt"]
    band.line.fill.background()

    band_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.0), Inches(4.6), Inches(0.2))
    p = band_box.text_frame.paragraphs[0]
    p.text = "需求分析与系统设计汇报"
    set_text_style(p, 11, COLORS["primary_dark"], bold=True)

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.65), Inches(8.6), Inches(1.4))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "膳食营养分析系统"
    set_text_style(p, 28, COLORS["text"], bold=True)
    p = tf.add_paragraph()
    p.text = "DietApp Requirements & System Design"
    set_text_style(p, 16, COLORS["muted"])

    desc_box = slide.shapes.add_textbox(Inches(0.78), Inches(3.0), Inches(5.8), Inches(1.0))
    tf = desc_box.text_frame
    p = tf.paragraphs[0]
    p.text = "基于 Kotlin、Jetpack Compose、Retrofit、AI 分析与图像识别的移动端健康管理应用。"
    set_text_style(p, 14, COLORS["text"])

    meta_box = slide.shapes.add_textbox(Inches(0.78), Inches(4.2), Inches(5.0), Inches(0.5))
    tf = meta_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"生成日期：{datetime.now().strftime('%Y-%m-%d')}"
    set_text_style(p, 11, COLORS["muted"])

    hero = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.25), Inches(3.75), Inches(4.85)
    )
    hero.fill.solid()
    hero.fill.fore_color.rgb = COLORS["surface"]
    hero.line.color.rgb = COLORS["border"]

    add_tag(slide, "移动健康", Inches(8.9), Inches(1.6), COLORS["surface_alt"])
    add_tag(slide, "AI 辅助", Inches(10.45), Inches(1.6), COLORS["surface_alt"])

    for idx, (title, color) in enumerate(
        [("AI 饮食分析", COLORS["primary"]), ("饮食日记", COLORS["secondary"]), ("数据看板", COLORS["teal"]), ("个人设置", COLORS["primary_dark"])]
    ):
        top = Inches(2.15 + idx * 0.86)
        chip = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.95), top, Inches(2.9), Inches(0.58)
        )
        chip.fill.solid()
        chip.fill.fore_color.rgb = color
        chip.line.fill.background()
        box = slide.shapes.add_textbox(Inches(9.2), top + Inches(0.12), Inches(2.2), Inches(0.25))
        p = box.text_frame.paragraphs[0]
        p.text = title
        set_text_style(p, 13, COLORS["white"], bold=True)

    add_footer(slide, 1)


def build_overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "项目概述", "围绕日常饮食记录、营养分析与健康反馈构建一体化体验")

    add_bullet_block(
        slide,
        Inches(0.75),
        Inches(1.7),
        Inches(4.0),
        Inches(4.9),
        "项目目标",
        [
            "提供低门槛的饮食录入方式：文字、拍照识别、食物选择器。",
            "借助 AI 与本地知识库生成结构化营养分析结果。",
            "将饮食、运动、体重趋势整合到一个移动端产品中。",
            "帮助用户形成可持续的健康记录与复盘习惯。",
        ],
    )
    add_bullet_block(
        slide,
        Inches(4.95),
        Inches(1.7),
        Inches(4.0),
        Inches(4.9),
        "核心痛点",
        [
            "手工估算热量和三大营养素成本高、误差大。",
            "用户很难把单餐记录与长期趋势联系起来。",
            "传统健康工具常见“记录分散、反馈滞后、建议不具体”。",
            "需要兼顾分析准确性、交互效率与陪伴式体验。",
        ],
    )
    add_footer(slide, 2)


def build_requirements_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "需求分析", "功能需求与非功能需求并重")

    add_bullet_block(
        slide,
        Inches(0.75),
        Inches(1.7),
        Inches(4.0),
        Inches(4.9),
        "功能需求",
        [
            "AI 助手：文本输入、拍照识别、相册选择、食物选择器录入。",
            "饮食日记：按日期查看记录、展示餐次、删除食物条目。",
            "数据看板：热量趋势、营养结构、体重变化、教练提示。",
            "个人中心：身高、每日热量目标保存与读取。",
        ],
    )
    add_bullet_block(
        slide,
        Inches(4.95),
        Inches(1.7),
        Inches(4.0),
        Inches(4.9),
        "非功能需求",
        [
            "性能：协程异步网络请求，保证主线程交互流畅。",
            "可靠性：Repository 统一封装接口并做好异常处理。",
            "安全性：权限最小化、接口可控、配置数据本地持久化。",
            "易用性：底部导航清晰、信息分层明确、反馈及时。",
        ],
    )
    add_footer(slide, 3)


def build_architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "系统总体架构", "移动端 + 接口服务 + AI 分析 + 本地知识库协同工作")

    blocks = [
        ("Android 客户端", "Kotlin / Fragment + Compose\nMainActivity + 4 个核心页面", COLORS["primary"]),
        ("Repository / API 层", "DietRepository\nRetrofit + Gson + OkHttp", COLORS["secondary"]),
        ("智能与数据服务", "DeepSeek 对话分析\n百度食物识别\nRESTful 业务接口", COLORS["teal"]),
        ("本地辅助能力", "food_knowledge_base.json\nSharedPreferences\n图表展示与状态缓存", COLORS["primary_dark"]),
    ]

    lefts = [Inches(0.8), Inches(3.35), Inches(5.9), Inches(8.45)]
    for idx, (title, body, accent) in enumerate(blocks):
        add_card(slide, lefts[idx], Inches(2.15), Inches(2.15), Inches(2.65), title, body, accent)
        if idx < len(blocks) - 1:
            arrow = slide.shapes.add_textbox(lefts[idx] + Inches(2.17), Inches(3.0), Inches(0.35), Inches(0.3))
            p = arrow.text_frame.paragraphs[0]
            p.text = "→"
            set_text_style(p, 24, COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)

    note = slide.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(10.8), Inches(0.9))
    tf = note.text_frame
    p = tf.paragraphs[0]
    p.text = "架构特点：前端重交互，Repository 负责数据访问抽象，AI 服务负责生成分析结果，本地知识库用于增强解释质量。"
    set_text_style(p, 13, COLORS["text"])
    add_footer(slide, 4)


def build_modules_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "系统功能模块", "四大模块围绕“记录 - 分析 - 保存 - 复盘”形成闭环")

    cards = [
        ("AI 饮食分析模块", "接收文字、拍照、相册和选择器输入；调用 RAG 与 DeepSeek 输出食物明细、热量与营养建议。", COLORS["primary"]),
        ("饮食日记模块", "按日期展示记录；按餐次聚合食物；支持查看总热量与删除单项食物。", COLORS["secondary"]),
        ("数据看板模块", "展示今日摘要、热量趋势、营养占比、体重变化和 AI 教练提示。", COLORS["teal"]),
        ("个人设置模块", "保存身高与每日热量目标，为分析和运动消耗估算提供个性化参数。", COLORS["primary_dark"]),
    ]

    positions = [
        (Inches(0.8), Inches(1.9)),
        (Inches(6.6), Inches(1.9)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.6), Inches(4.2)),
    ]
    for (title, body, accent), (left, top) in zip(cards, positions):
        add_card(slide, left, top, Inches(5.1), Inches(1.8), title, body, accent)
    add_footer(slide, 5)


def build_entities_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "核心数据实体", "数据模型支撑页面展示、统计分析与接口交互")

    entities = [
        ("FoodItem", "食物名称、份量、单位、重量、热量、蛋白质、碳水、脂肪"),
        ("DietRecord", "日期、餐次、食物列表、总热量、总营养素、AI 建议"),
        ("ExerciseRecord", "日期、运动类型、时长、消耗热量"),
        ("DailySummary", "当日摄入、运动消耗、净热量聚合"),
        ("BodyTrend", "体重、体脂、BMI 的时间序列变化"),
        ("UserPreferences", "身高、每日热量目标、本地个性化配置"),
    ]

    for idx, (name, desc) in enumerate(entities):
        top = Inches(1.85 + idx * 0.78)
        tag = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), top, Inches(1.75), Inches(0.52)
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLORS["surface_alt"]
        tag.line.fill.background()
        box = slide.shapes.add_textbox(Inches(1.08), top + Inches(0.1), Inches(1.4), Inches(0.2))
        p = box.text_frame.paragraphs[0]
        p.text = name
        set_text_style(p, 12, COLORS["primary_dark"], bold=True)

        body = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.8), top, Inches(8.55), Inches(0.52)
        )
        body.fill.solid()
        body.fill.fore_color.rgb = COLORS["surface"]
        body.line.color.rgb = COLORS["border"]
        box = slide.shapes.add_textbox(Inches(3.0), top + Inches(0.09), Inches(8.1), Inches(0.25))
        p = box.text_frame.paragraphs[0]
        p.text = desc
        set_text_style(p, 11, COLORS["text"])

    add_footer(slide, 6)


def build_pages_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "页面结构与导航", "以底部导航为中心组织用户高频操作")

    center = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.75), Inches(2.8), Inches(2.0), Inches(1.0)
    )
    center.fill.solid()
    center.fill.fore_color.rgb = COLORS["primary"]
    center.line.fill.background()
    box = slide.shapes.add_textbox(Inches(4.95), Inches(3.08), Inches(1.6), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = "Bottom Navigation"
    set_text_style(p, 15, COLORS["white"], bold=True, align=PP_ALIGN.CENTER)

    items = [
        ("AI 助手", "文本/拍照输入\n结果编辑与保存", Inches(1.0), Inches(1.5), COLORS["surface_alt"]),
        ("饮食日记", "日期切换\n餐次记录查看", Inches(8.8), Inches(1.5), COLORS["surface"]),
        ("数据看板", "摘要卡\n图表可视化", Inches(1.0), Inches(4.5), COLORS["surface"]),
        ("个人中心", "目标设置\n健康建议", Inches(8.8), Inches(4.5), COLORS["surface_alt"]),
    ]
    for title, body, left, top, fill in items:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, Inches(2.7), Inches(1.35)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = COLORS["border"]
        tbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.18), Inches(2.3), Inches(0.9))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        set_text_style(p, 15, COLORS["text"], bold=True)
        p = tf.add_paragraph()
        p.text = body
        set_text_style(p, 11, COLORS["muted"])

    arrows = [
        (Inches(3.75), Inches(2.2), "↘"),
        (Inches(7.3), Inches(2.2), "↙"),
        (Inches(3.75), Inches(4.25), "↗"),
        (Inches(7.3), Inches(4.25), "↖"),
    ]
    for left, top, text in arrows:
        tbox = slide.shapes.add_textbox(left, top, Inches(0.4), Inches(0.4))
        p = tbox.text_frame.paragraphs[0]
        p.text = text
        set_text_style(p, 22, COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)

    add_footer(slide, 7)


def build_stack_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "关键技术栈", "围绕 Android 原生开发与外部智能能力构建")

    tech_cards = [
        ("前端框架", ["Kotlin", "Jetpack Compose", "Fragment + Navigation"]),
        ("数据访问", ["Retrofit", "OkHttp", "Gson", "Kotlin Coroutines"]),
        ("能力组件", ["CameraX", "MPAndroidChart", "SharedPreferences"]),
        ("智能服务", ["DeepSeek Chat", "百度食物识别", "本地 RAG 知识库"]),
    ]

    for idx, (title, items) in enumerate(tech_cards):
        left = Inches(0.8 + idx * 3.0)
        add_card(slide, left, Inches(2.1), Inches(2.6), Inches(3.2), title, "\n".join(items), COLORS["primary"] if idx % 2 == 0 else COLORS["secondary"])

    note = slide.shapes.add_textbox(Inches(0.85), Inches(5.8), Inches(10.7), Inches(0.7))
    p = note.text_frame.paragraphs[0]
    p.text = "技术选型目标：确保移动端交互流畅、接口调用清晰、图表展示直观，并为后续扩展后台管理与个性化推荐留出空间。"
    set_text_style(p, 13, COLORS["text"])
    add_footer(slide, 8)


def build_highlights_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "系统亮点", "在传统饮食记录基础上增强智能化与可解释性")

    highlights = [
        ("多输入方式", "文本、拍照、相册与食物选择器并行，适配不同使用场景。"),
        ("RAG + AI 联合分析", "本地食物知识库与大模型结合，提升结果解释性。"),
        ("可编辑后再保存", "AI 结果不是黑盒，用户可以逐项修正后写入日记。"),
        ("趋势驱动复盘", "把单次记录延伸到周/月趋势，帮助形成长期习惯。"),
    ]
    for idx, (title, body) in enumerate(highlights):
        add_card(
            slide,
            Inches(0.9 + (idx % 2) * 5.5),
            Inches(2.0 + (idx // 2) * 2.0),
            Inches(4.9),
            Inches(1.55),
            title,
            body,
            [COLORS["primary"], COLORS["secondary"], COLORS["teal"], COLORS["primary_dark"]][idx],
        )
    add_footer(slide, 9)


def build_summary_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_canvas(slide)
    add_header(slide, "总结与后续工作", "当前系统已形成可运行的移动端原型，后续可继续深化")

    add_bullet_block(
        slide,
        Inches(0.85),
        Inches(1.9),
        Inches(4.9),
        Inches(4.6),
        "当前成果",
        [
            "完成 AI 饮食分析、饮食日记、数据看板和个人设置四大功能。",
            "打通移动端导航、接口访问、图像识别与可视化展示流程。",
            "形成较完整的需求分析与系统设计文档及汇报材料。",
        ],
    )
    add_bullet_block(
        slide,
        Inches(6.0),
        Inches(1.9),
        Inches(4.9),
        Inches(4.6),
        "后续优化方向",
        [
            "补充后台管理端与用户账户体系。",
            "完善异常提示、测试覆盖与数据校验逻辑。",
            "增强个性化建议、健康目标跟踪和数据导出能力。",
            "优化界面视觉一致性与实际部署方案。",
        ],
    )

    thanks = slide.shapes.add_textbox(Inches(0.88), Inches(6.7), Inches(10.2), Inches(0.35))
    p = thanks.text_frame.paragraphs[0]
    p.text = "谢谢观看"
    set_text_style(p, 20, COLORS["primary_dark"], bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 10)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs)
    build_overview_slide(prs)
    build_requirements_slide(prs)
    build_architecture_slide(prs)
    build_modules_slide(prs)
    build_entities_slide(prs)
    build_pages_slide(prs)
    build_stack_slide(prs)
    build_highlights_slide(prs)
    build_summary_slide(prs)

    prs.core_properties.title = "膳食营养分析系统汇报PPT"
    prs.core_properties.subject = "需求分析与系统设计"
    prs.core_properties.comments = "使用 python-pptx 自动生成"
    prs.save(str(OUTPUT_PATH))
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
